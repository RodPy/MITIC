#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador de Actas de Recepción Provisoria (Bienes y Solución) para los 40 laboratorios TIC.
LPI MITIC/PAAD UEP N° 02/2025 – ID DNCP 468.772 – Contrato N° 122/2026 – TECHNOMA S.A.E.C.A.

Toma como base las dos actas modelo finales (Word, Colegio Nacional Julio Correa) y, para cada
establecimiento del CSV, genera:
  - ACTA_RECEPCION_PROVISORIA_BIENES_<código>_<colegio>.docx / .pdf
  - ACTA_RECEPCION_PROVISORIA_SOLUCION_<código>_<colegio>.docx / .pdf
Los PDF salen completables (casillas, cuadro de observaciones y campos de firma digital).
Al firmar el Fiscal MITIC se traban las casillas y el cuadro de observaciones; las firmas del
Proveedor y del Director/a siguen habilitadas.
También genera reporte_validacion.csv con los datos faltantes o dudosos de cada colegio.

Requisitos:
    pip install pandas python-docx pymupdf
    + LibreOffice instalado (recomendado)  o  Microsoft Word + "pip install docx2pdf"

Uso:
    python generar_actas.py --csv LABS_ALL_SCHOOLS_FULL_ITEMS.csv \
        --plantilla-bienes ACTA_RECEPCION_PROVISORIA_BIENES_1109202_Colegio_Nacional_Julio_Correa.docx \
        --plantilla-solucion ACTA_RECEPCION_PROVISORIA_SOLUCION_1109202_Colegio_Nacional_Julio_Correa.docx \
        --salida actas_generadas

Opciones:
    --codigo 1109202 1001025   solo esos establecimientos
    --sin-pdf                  solo genera los Word
    --motor-pdf libreoffice|word|auto
    --agregar-faltantes        agrega los kits estándar del Lote 1 que no figuren en el CSV
                               (Herramientas complementarias, Herramientas de mano, Insumos)
"""
import argparse, copy, os, re, shutil, subprocess, sys, tempfile, unicodedata
from pathlib import Path

import pandas as pd
import docx
from docx.table import _Row

# --------------------------------------------------------------------------------------
# Valores que tienen las plantillas (Colegio Nacional Julio Correa) y que se reemplazan
# --------------------------------------------------------------------------------------
PLANTILLA = {
    'codigo': '1109202',
    'colegio_mayus': 'COLEGIO NACIONAL JULIO CORREA',
    'colegio_titulo': 'Colegio Nacional Julio Correa',
    'ciudad_depto': 'Luque, Departamento Central',
    'ciudad': 'Luque',
    'depto': 'Central',
    'fecha_texto': '12 días del mes de octubre de 2026',
    'fecha': '12/10/2026',
    'serie_aire': '54004157101A3010250481',
}

# --------------------------------------------------------------------------------------
# Catálogo: nombre en el acta, orden, precio unitario (Gs., según contrato) y lote
# --------------------------------------------------------------------------------------
CATALOGO = [
    # clave CSV,                                      nombre en el acta,                               precio,    lote, requiere serie
    ('Notebook',                                      'Notebook',                                      6065175,   '1', True),
    ('Placa de programación',                         'Placa de programación',                         1679789,   '1', False),
    ('Kit de Robótica',                               'Kit de Robótica',                               2560472,   '1', False),
    ('Impresora 3D',                                  'Impresora 3D',                                  2493875,   '1', True),
    ('Herramientas complementarias electrónica (Kit)', 'Herramientas complementarias electrónica (Kit)', 11560575, '1', False),
    ('TV de 55"',                                     'TV de 55"',                                     6031450,   '1', True),
    ('Gabinete de carga para notebooks',              'Gabinete de carga para notebooks',              5108450,   '1', False),
    ('Enrutador Ethernet',                            'Enrutador Ethernet',                            9036524,   '1', True),
    ('Insumos/consumibles para equipamientos (Kit)',  'Insumos/consumibles para equipamientos (Kit)',  1993324,   '1', False),
    ('Herramientas de mano',                          'Herramientas de mano',                          4291950,   '1', False),
    ('SILLA',                                         'Silla',                                         434590,    '1', False),
    ('MESA',                                          'Mesa de trabajo',                               1182860,   '1', False),
    ('MESA RECTANGULAR',                              'Escritorio para facilitador',                   1327700,   '1', False),
    ('ARMARIO',                                       'Armario bajo con puertas corredizas',           1798998,   '1', False),
    ('Impresora 3D (resina)',                         'Impresora 3D (resina)',                         5388900,   '2', True),
    ('Cortadora Láser',                               'Cortadora Láser',                               4176575,   '2', True),
    ('Escáner 3D',                                    'Escáner 3D',                                    14150300,  '2', True),
    ('Kit de IA con Raspberry Pi 5',                  'Kit de IA con Raspberry Pi 5',                  4494300,   '2', False),
    ('Monitor 22" FHD',                               'Monitor 22" FHD',                               None,      '2', True),   # parte del ítem Escáner 3D
    ('PC',                                            'PC',                                            None,      '2', True),   # parte del ítem Escáner 3D
]
CAT = {c[0].upper(): c for c in CATALOGO}
ORDEN = {c[0].upper(): i for i, c in enumerate(CATALOGO)}
AIRE = 'AIRE ACONDICIONADO'

# Ítems estándar del Lote 1 que pueden faltar en el CSV (para --agregar-faltantes)
FALTANTES_ESTANDAR = {
    'Herramientas complementarias electrónica (Kit)': ('S/D', 'ANENG', 1),
    'Insumos/consumibles para equipamientos (Kit)':   ('S/D', 'GENERICO', 1),
    'Herramientas de mano':                           ('S/D', 'TOTAL', 1),
}
# Ítems que todo laboratorio debería tener según el contrato (para el reporte)
ESPERADOS = {'Notebook': 21, 'Placa de programación': 15, 'Kit de Robótica': 4, 'Impresora 3D': 1,
             'Herramientas complementarias electrónica (Kit)': 1, 'TV de 55"': 1, 'Gabinete de carga para notebooks': 1,
             'Enrutador Ethernet': 1, 'Insumos/consumibles para equipamientos (Kit)': 1, 'Herramientas de mano': 1,
             'SILLA': 21, 'MESA': 10, 'MESA RECTANGULAR': 1, 'ARMARIO': 2, AIRE: 1,
             'Impresora 3D (resina)': 1, 'Cortadora Láser': 1, 'Escáner 3D': 1, 'Kit de IA con Raspberry Pi 5': 1,
             'Monitor 22" FHD': 1, 'PC': 1}

MESES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto',
         'septiembre', 'octubre', 'noviembre', 'diciembre']
MINUSCULAS = {'de', 'del', 'la', 'las', 'los', 'el', 'y', 'e'}


# --------------------------------------------------------------------------------------
# Utilidades de texto
# --------------------------------------------------------------------------------------
def limpio(x):
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return ''
    return ' '.join(str(x).split())


def titulo(nombre):
    """'COLEGIO NACIONAL KA'ARENDY' -> 'Colegio Nacional Ka'arendy' (respeta números romanos y siglas)."""
    out = []
    for i, w in enumerate(limpio(nombre).split(' ')):
        partes = []
        for p in w.split('-'):
            low = p.lower()
            if re.fullmatch(r'[IVXL]+', p):                 # II, IV...
                partes.append(p)
            elif i > 0 and low in MINUSCULAS:
                partes.append(low)
            elif re.fullmatch(r'([A-ZÁÉÍÓÚÑ]\.)+', p):      # E.M.D.
                partes.append(p)
            else:
                partes.append(low[:1].upper() + low[1:])
        out.append('-'.join(partes))
    return ' '.join(out)


def slug(s):
    s = unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode()
    return re.sub(r'[^A-Za-z0-9]+', '_', s).strip('_')


def fecha_texto(fecha):
    d, m, y = [int(x) for x in fecha.split('/')]
    return f'{d} {"día" if d == 1 else "días"} del mes de {MESES[m - 1]} de {y}'


def guaranies(n):
    return '-' if n is None else 'Gs. ' + f'{n:,}'.replace(',', '.')


# --------------------------------------------------------------------------------------
# Utilidades Word
# --------------------------------------------------------------------------------------
def reemplazar_en_parrafo(p, viejo, nuevo):
    """Reemplaza texto aunque esté repartido en varios 'runs', conservando el formato del run
    que contiene la mayor parte del texto buscado."""
    if viejo not in p.text:
        return 0
    n = 0
    while True:
        runs = p.runs
        texto = ''.join(r.text for r in runs)
        i = texto.find(viejo)
        if i < 0:
            return n
        j = i + len(viejo)
        pos = 0
        afectados = []          # (run, inicio, fin) del tramo dentro de cada run
        for r in runs:
            a, b = pos, pos + len(r.text)
            ini, fin = max(a, i), min(b, j)
            if ini < fin:
                afectados.append((r, ini - a, fin - a))
            pos = b
        principal = max(afectados, key=lambda t: t[2] - t[1])[0]
        for r, ini, fin in afectados:
            t = r.text
            r.text = t[:ini] + (nuevo if r is principal else '') + t[fin:]
        n += 1


def reemplazar_en_documento(d, pares):
    def parrafos():
        for p in d.paragraphs:
            yield p
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    yield from c.paragraphs
    for p in parrafos():
        for viejo, nuevo in pares:
            if viejo != nuevo:
                reemplazar_en_parrafo(p, viejo, nuevo)


def poner_texto_celda(celda, texto):
    p = celda.paragraphs[0]
    for extra in celda.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    if p.runs:
        p.runs[0].text = texto
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run(texto)


def tabla_con_encabezado(d, primer_titulo, contiene):
    for t in d.tables:
        h = [c.text.strip() for c in t.rows[0].cells]
        if h and h[0] == primer_titulo and contiene in h:
            return t
    raise ValueError(f'No se encontró en la plantilla la tabla con la columna "{contiene}"')


def rellenar_tabla(t, filas):
    """Usa la 1.ª fila de datos como modelo de formato, borra las demás y agrega las nuevas."""
    modelo = copy.deepcopy(t.rows[1]._tr)
    for row in list(t.rows)[1:]:
        t._tbl.remove(row._tr)
    for valores in filas:
        tr = copy.deepcopy(modelo)
        t._tbl.append(tr)
        row = _Row(tr, t)
        for celda, v in zip(row.cells, valores):
            poner_texto_celda(celda, str(v))


# --------------------------------------------------------------------------------------
# Datos
# --------------------------------------------------------------------------------------
def leer_csv(ruta):
    for enc in ('utf-8-sig', 'latin-1'):
        try:
            df = pd.read_csv(ruta, dtype=str, encoding=enc)
            break
        except UnicodeDecodeError:
            continue
    df.columns = [c.strip() for c in df.columns]
    necesarias = ['N°', 'Descripción del bien', 'Cantidad', 'Modelo', 'Marca', 'Número de Serial',
                  'ID del llamado', 'Empresa', 'Contrato N°', 'Colegio', 'Departamento',
                  'Codigo de establecimiento', 'Ciudad', 'Fecha Programada']
    faltan = [c for c in necesarias if c not in df.columns]
    if faltan:
        sys.exit(f'ERROR: al CSV le faltan las columnas: {faltan}')
    for c in df.columns:
        df[c] = df[c].map(limpio)
    df['clave'] = df['Descripción del bien'].str.upper()
    df['cant'] = pd.to_numeric(df['Cantidad'], errors='coerce').fillna(1).astype(int)
    df['n'] = pd.to_numeric(df['N°'], errors='coerce')
    return df


def items_establecimiento(g, agregar_faltantes):
    g = g.copy()
    desconocidos = sorted(set(g['clave']) - set(CAT) - {AIRE})
    if agregar_faltantes:
        base = g.iloc[0]
        for clave, (modelo, marca, cant) in FALTANTES_ESTANDAR.items():
            if clave.upper() not in set(g['clave']):
                fila = base.copy()
                fila['Descripción del bien'] = clave; fila['clave'] = clave.upper()
                fila['Modelo'] = modelo; fila['Marca'] = marca; fila['Número de Serial'] = ''
                fila['cant'] = cant; fila['n'] = 99999
                g = pd.concat([g, fila.to_frame().T], ignore_index=True)
    g['orden'] = g['clave'].map(ORDEN)
    bienes = g[g['clave'].isin(CAT)].sort_values(['orden', 'n'])
    aire = g[g['clave'] == AIRE]
    return bienes, aire, desconocidos


def validar(codigo, info, g, bienes, aire, desconocidos):
    obs = []
    tot = g.groupby('clave')['cant'].sum()
    for clave, esperado in ESPERADOS.items():
        real = int(tot.get(clave.upper(), 0))
        if real != esperado:
            obs.append(f'{clave}: {real} en el CSV (contrato: {esperado})')
    for _, r in pd.concat([bienes, aire]).iterrows():
        c = CAT.get(r['clave'], (None, None, None, None, True))
        if c[4] and r['Número de Serial'] in ('', 'N/A', 'nan'):
            obs.append(f'{r["Descripción del bien"]} (N° {r["N°"]}) sin número de serie')
        if '¡' in r['Número de Serial']:
            obs.append(f'{r["Descripción del bien"]}: serie con carácter dudoso "{r["Número de Serial"]}"')
    if desconocidos:
        obs.append(f'Descripciones no reconocidas (no se incluyeron): {desconocidos}')
    return obs


# --------------------------------------------------------------------------------------
# Generación de los Word
# --------------------------------------------------------------------------------------
def pares_reemplazo(info, serie_aire):
    return [
        (PLANTILLA['ciudad_depto'], f'{info["ciudad"]}, Departamento {info["depto"]}'),
        (PLANTILLA['fecha_texto'], fecha_texto(info['fecha'])),
        (PLANTILLA['colegio_mayus'], info['colegio'].upper()),
        (PLANTILLA['colegio_titulo'], titulo(info['colegio'])),
        (PLANTILLA['codigo'], info['codigo']),
        (PLANTILLA['fecha'], info['fecha']),
        (PLANTILLA['serie_aire'], serie_aire),
    ]


def tabla_establecimiento(d, info):
    t = tabla_con_encabezado(d, 'Código de Establecimiento', 'Fecha SAT')
    for celda, v in zip(t.rows[1].cells, [info['codigo'], info['colegio'].upper(), info['depto'], info['ciudad'], info['fecha']]):
        poner_texto_celda(celda, v)


def generar_bienes(plantilla, destino, info, bienes, serie_aire):
    d = docx.Document(plantilla)
    tabla_establecimiento(d, info)
    filas = []
    for i, (_, r) in enumerate(bienes.iterrows(), 1):
        _, nombre, precio, lote, _ = CAT[r['clave']]
        desc = nombre + (f' ({r["cant"]} unid.)' if r['cant'] > 1 else '')
        filas.append([i, desc, r['Modelo'] or '-', r['Marca'] or '-', r['Número de Serial'] or 'N/A', guaranies(precio),
                      r['ID del llamado'] or '468772', r['Empresa'] or 'TECHNOMA S.A.E.C.A.', r['Contrato N°'] or '122/2026', lote])
    rellenar_tabla(tabla_con_encabezado(d, 'N°', 'Número de Serial'), filas)
    reemplazar_en_documento(d, pares_reemplazo(info, serie_aire))
    d.core_properties.title = f'Acta de Recepción Provisoria de los Bienes – {info["codigo"]}'
    d.save(destino)


def generar_solucion(plantilla, destino, info, bienes, serie_aire):
    d = docx.Document(plantilla)
    tabla_establecimiento(d, info)
    filas = []
    agrupado = bienes.groupby('clave', sort=False).agg(cant=('cant', 'sum'), marca=('Marca', 'first'), modelo=('Modelo', 'first'))
    for i, (clave, r) in enumerate(agrupado.iterrows(), 1):
        _, nombre, _, lote, _ = CAT[clave]
        mm = r.marca if r.modelo in ('', 'S/D') else f'{r.marca} – {r.modelo}'
        filas.append([i, nombre, mm, int(r.cant), lote])
    rellenar_tabla(tabla_con_encabezado(d, 'N°', 'Marca – Modelo'), filas)
    reemplazar_en_documento(d, pares_reemplazo(info, serie_aire))
    d.core_properties.title = f'Acta de Recepción Provisoria de la Solución – {info["codigo"]}'
    d.save(destino)


def restos_plantilla(ruta, info):
    """Devuelve textos de Julio Correa que hayan quedado sin reemplazar."""
    if info['codigo'] == PLANTILLA['codigo']:
        return []
    d = docx.Document(ruta)
    textos = [p.text for p in d.paragraphs] + [c.text for t in d.tables for row in t.rows for c in row.cells]
    full = '\n'.join(textos)
    return [v for k, v in PLANTILLA.items() if k not in ('depto', 'fecha', 'fecha_texto') and v in full]


# --------------------------------------------------------------------------------------
# PDF y campos completables
# --------------------------------------------------------------------------------------
def buscar_soffice():
    for nombre in ('soffice', 'libreoffice'):
        p = shutil.which(nombre)
        if p:
            return p
    for p in [r'C:\Program Files\LibreOffice\program\soffice.exe',
              r'C:\Program Files (x86)\LibreOffice\program\soffice.exe',
              '/Applications/LibreOffice.app/Contents/MacOS/soffice']:
        if os.path.exists(p):
            return p
    return None


def a_pdf(docx_path, motor):
    pdf = Path(docx_path).with_suffix('.pdf')
    soffice = buscar_soffice()
    if motor in ('libreoffice', 'auto') and soffice:
        perfil = Path(tempfile.gettempdir()) / 'lo_perfil_actas'
        subprocess.run([soffice, f'-env:UserInstallation=file:///{perfil.as_posix().lstrip("/")}', '--headless',
                        '--convert-to', 'pdf', '--outdir', str(pdf.parent), str(docx_path)],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        return pdf
    if motor in ('word', 'auto'):
        try:
            from docx2pdf import convert
        except ImportError:
            sys.exit('ERROR: no se encontró LibreOffice ni docx2pdf. Instale LibreOffice o ejecute: pip install docx2pdf')
        convert(str(docx_path), str(pdf))
        return pdf
    sys.exit('ERROR: no se encontró LibreOffice (soffice). Instálelo o use --motor-pdf word')


def agregar_campos(pdf_path):
    import pymupdf as fitz
    tmp = str(pdf_path) + '.tmp'
    d = fitz.open(pdf_path)
    usados = set(); en_tabla = [0]

    def unico(n):
        k, i = n, 2
        while k in usados:
            k = f'{n}_{i}'; i += 1
        usados.add(k); return k

    for p in d:
        palabras = p.get_text('words')
        # Casillas
        for b in p.search_for('☐'):
            cy = (b.y0 + b.y1) / 2
            linea = sorted(w for w in palabras if abs((w[1] + w[3]) / 2 - cy) < 4 and w[0] > b.x1 - 1)
            etiqueta = ' '.join(w[4] for w in linea[:3])
            r = fitz.Rect(b.x0 - 0.5, cy - 4.6, b.x0 + 8.7, cy + 4.6)
            p.draw_rect(r, color=None, fill=(1, 1, 1), overlay=True)
            if b.x0 > p.rect.width / 2:                       # columna Verificación/Resultado
                k = en_tabla[0]; en_tabla[0] += 1
                sec, fila = ('Readecuacion', k // 2 + 1) if k < 6 else ('Funcionamiento', (k - 6) // 2 + 1)
                nombre = f'{sec}_{fila}_' + ('No_conforme' if k % 2 else 'Conforme')
            else:
                nombre = 'Sin_observaciones' if etiqueta.lower().startswith('sin') else 'Con_observaciones'
            w = fitz.Widget(); w.field_type = fitz.PDF_WIDGET_TYPE_CHECKBOX; w.rect = r
            w.field_name = unico(nombre); w.border_color = (0, 0, 0); w.border_width = 0.8; w.fill_color = (1, 1, 1)
            w.field_value = False
            p.add_widget(w)
        # Cuadro de observaciones (encabezado centrado de la tabla)
        enc = [h for h in p.search_for('Observaciones  y/o pendientes') + p.search_for('Observaciones y/o pendientes')
               if h.x0 > p.rect.width * 0.3]
        if enc:
            h = enc[0]
            debajo = [w for w in palabras if w[1] > h.y1 + 8]
            fondo = min(w[1] for w in debajo) - 14 if debajo else h.y1 + 120
            w = fitz.Widget(); w.field_type = fitz.PDF_WIDGET_TYPE_TEXT
            w.rect = fitz.Rect(78, h.y1 + 4, p.rect.width - 78, fondo)
            w.field_name = unico('Observaciones_pendientes'); w.field_flags = fitz.PDF_TX_FIELD_IS_MULTILINE
            w.text_font = 'Helv'; w.text_fontsize = 10; w.border_width = 0
            p.add_widget(w)
        # Firmas (sobre cada línea de firma)
        lineas = sorted([w for w in palabras if w[4].startswith('_____')], key=lambda w: w[0])
        roles = ['Fiscal_MITIC', 'Representante_Proveedor', 'Director_establecimiento']
        for i, u in enumerate(lineas):
            w = fitz.Widget(); w.field_type = fitz.PDF_WIDGET_TYPE_SIGNATURE
            w.rect = fitz.Rect(u[0], u[1] - 38, u[2], u[1] + 2)
            w.field_name = unico('Firma_' + (roles[i] if i < len(roles) else str(i + 1))); w.border_width = 0
            p.add_widget(w)
    # Campos que se completan (casillas y observaciones): los traba la firma del Fiscal
    completables = [w.field_name for p in d for w in p.widgets()
                    if w.field_type != fitz.PDF_WIDGET_TYPE_SIGNATURE]
    lista = ' '.join('(' + n.replace('(', '\\(').replace(')', '\\)') + ')' for n in completables)
    for p in d:
        for w in p.widgets():
            if w.field_type == fitz.PDF_WIDGET_TYPE_SIGNATURE:
                # Aspecto en blanco (sin el sello "SIGN" de PyMuPDF)
                ancho, alto = w.rect.width, w.rect.height
                ap = d.get_new_xref()
                d.update_object(ap, f'<</Type/XObject/Subtype/Form/BBox[0 0 {ancho:.2f} {alto:.2f}]/Resources<<>>>>')
                d.update_stream(ap, b' ')
                d.xref_set_key(w.xref, 'AP', f'<</N {ap} 0 R>>')
                # La firma del Fiscal traba solo los campos completables; las demás firmas quedan libres
                if w.field_name == 'Firma_Fiscal_MITIC' and completables:
                    d.xref_set_key(w.xref, 'Lock', f'<</Type/SigFieldLock/Action/Include/Fields[{lista}]>>')
    d.save(tmp, garbage=3, deflate=True)
    d.close()
    os.replace(tmp, pdf_path)


# --------------------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description='Genera las actas de recepción provisoria para cada laboratorio.')
    ap.add_argument('--csv', required=True)
    ap.add_argument('--plantilla-bienes', required=True)
    ap.add_argument('--plantilla-solucion', required=True)
    ap.add_argument('--salida', default='actas_generadas')
    ap.add_argument('--codigo', nargs='*', help='Generar solo estos códigos de establecimiento')
    ap.add_argument('--sin-pdf', action='store_true')
    ap.add_argument('--motor-pdf', choices=['auto', 'libreoffice', 'word'], default='auto')
    ap.add_argument('--agregar-faltantes', action='store_true')
    a = ap.parse_args()

    for plantilla in (a.plantilla_bienes, a.plantilla_solucion):
        texto = '\n'.join(p.text for p in docx.Document(plantilla).paragraphs)
        if PLANTILLA['codigo'] not in texto or PLANTILLA['fecha_texto'] not in texto:
            sys.exit(f'ERROR: la plantilla {plantilla} no contiene los datos del Colegio Nacional Julio Correa '
                     f'que el script reemplaza (código {PLANTILLA["codigo"]}, "{PLANTILLA["fecha_texto"]}").')

    df = leer_csv(a.csv)
    salida = Path(a.salida); salida.mkdir(parents=True, exist_ok=True)
    codigos = a.codigo or list(dict.fromkeys(df['Codigo de establecimiento']))
    reporte = []

    for n, codigo in enumerate(codigos, 1):
        g = df[df['Codigo de establecimiento'] == codigo]
        if g.empty:
            print(f'[{n}/{len(codigos)}] {codigo}: no está en el CSV, se omite'); continue
        f0 = g.iloc[0]
        info = {'codigo': codigo, 'colegio': f0['Colegio'], 'depto': f0['Departamento'],
                'ciudad': f0['Ciudad'], 'fecha': f0['Fecha Programada']}
        bienes, aire, desconocidos = items_establecimiento(g, a.agregar_faltantes)
        serie_aire = aire['Número de Serial'].iloc[0] if len(aire) and aire['Número de Serial'].iloc[0] else '[sin dato]'

        carpeta = salida / f'{codigo}_{slug(titulo(info["colegio"]))}'
        carpeta.mkdir(exist_ok=True)
        base = f'{codigo}_{slug(titulo(info["colegio"]))}'
        f_b = carpeta / f'ACTA_RECEPCION_PROVISORIA_BIENES_{base}.docx'
        f_s = carpeta / f'ACTA_RECEPCION_PROVISORIA_SOLUCION_{base}.docx'
        generar_bienes(a.plantilla_bienes, f_b, info, bienes, serie_aire)
        generar_solucion(a.plantilla_solucion, f_s, info, bienes, serie_aire)

        obs = validar(codigo, info, g, bienes, aire, desconocidos)
        for f in (f_b, f_s):
            restos = restos_plantilla(f, info)
            if restos:
                obs.append(f'{f.name}: quedaron textos de la plantilla sin reemplazar: {restos}')
        if not a.sin_pdf:
            for f in (f_b, f_s):
                agregar_campos(a_pdf(f, a.motor_pdf))

        reporte.append({'Código': codigo, 'Colegio': info['colegio'], 'Departamento': info['depto'],
                        'Ciudad': info['ciudad'], 'Fecha SAT': info['fecha'], 'Bienes en el acta': len(bienes),
                        'Observaciones de datos': ' | '.join(obs) or 'OK'})
        print(f'[{n}/{len(codigos)}] {codigo} {info["colegio"]}: listo' + (f' ({len(obs)} observaciones)' if obs else ''))

    pd.DataFrame(reporte).to_csv(salida / 'reporte_validacion.csv', index=False, encoding='utf-8-sig')
    print(f'\nActas generadas en: {salida.resolve()}\nRevise el reporte: {(salida / "reporte_validacion.csv").resolve()}')


if __name__ == '__main__':
    main()
